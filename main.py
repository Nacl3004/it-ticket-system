from fastapi import FastAPI, HTTPException, Depends, WebSocket, WebSocketDisconnect, UploadFile, File, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import RedirectResponse
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
import pymysql
import os
from dotenv import load_dotenv
from datetime import datetime, timedelta
import jwt
import bcrypt
import json
import openpyxl
from io import BytesIO
import httpx
import hashlib
import hmac
import base64
import time
import random
import string
import uuid

load_dotenv()

os.environ["TZ"] = "Asia/Shanghai"
if hasattr(time, "tzset"):
    time.tzset()

app = FastAPI(title="IT运维工单系统")

# CORS配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def require_env(name: str) -> str:
    value = os.getenv(name)
    if not value:
        raise RuntimeError(f"Missing required environment variable: {name}")
    return value

# JWT配置
SECRET_KEY = require_env("SECRET_KEY")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 1440

# 数据库配置
DB_CONFIG = {
    "host": require_env("DB_HOST"),
    "port": int(os.getenv("DB_PORT", 3306)),
    "user": require_env("DB_USER"),
    "password": require_env("DB_PASSWORD"),
    "database": require_env("DB_NAME"),
    "charset": "utf8mb4"
}

DB_TIME_ZONE = "+08:00"

DEFAULT_WEBHOOK_TEMPLATES = {
    "ticket_created": {
        "name": "默认-工单创建",
        "title": "新工单提醒",
        "content": "**工单编号**: {ticket_no}\n\n**标题**: {title}\n\n**提交人**: {submitter}\n\n**优先级**: {priority}\n\n**时间**: {time}"
    },
    "ticket_claimed": {
        "name": "默认-工单认领",
        "title": "工单已被认领",
        "content": "**工单编号**: {ticket_no}\n\n**标题**: {title}\n\n**认领人**: {operator}\n\n**认领时间**: {time}"
    },
    "ticket_processing": {
        "name": "默认-工单处理中",
        "title": "工单处理中",
        "content": "**工单编号**: {ticket_no}\n\n**标题**: {title}\n\n**处理人**: {operator}\n\n**备注**: {message}\n\n**开始时间**: {time}"
    },
    "ticket_completed": {
        "name": "默认-工单完成",
        "title": "工单已完成",
        "content": "**工单编号**: {ticket_no}\n\n**标题**: {title}\n\n**处理人**: {operator}\n\n**解决方案**: {solution}\n\n**完成时间**: {time}"
    },
    "ticket_closed": {
        "name": "默认-工单结单",
        "title": "工单已结单",
        "content": "**工单编号**: {ticket_no}\n\n**标题**: {title}\n\n**结单人**: {operator}\n\n**结单时间**: {time}"
    }
}

# WebSocket连接管理
class ConnectionManager:
    def __init__(self):
        self.active_connections: dict = {}

    async def connect(self, websocket: WebSocket, user_id: int):
        await websocket.accept()
        self.active_connections[user_id] = websocket

    def disconnect(self, user_id: int):
        if user_id in self.active_connections:
            del self.active_connections[user_id]

    async def send_personal_message(self, message: dict, user_id: int):
        if user_id in self.active_connections:
            try:
                await self.active_connections[user_id].send_json(message)
            except:
                self.disconnect(user_id)

    async def broadcast(self, message: dict):
        for user_id, connection in list(self.active_connections.items()):
            try:
                await connection.send_json(message)
            except:
                self.disconnect(user_id)

manager = ConnectionManager()

# 数据库连接
def get_db():
    conn = pymysql.connect(**DB_CONFIG)
    try:
        if DB_TIME_ZONE:
            with conn.cursor() as cursor:
                cursor.execute("SET time_zone = %s", (DB_TIME_ZONE,))
        yield conn
    finally:
        conn.close()

# 数据库迁移
def check_and_migrate_db():
    conn = pymysql.connect(**DB_CONFIG)
    try:
        cursor = conn.cursor()
        if DB_TIME_ZONE:
            cursor.execute("SET time_zone = %s", (DB_TIME_ZONE,))
        
        # 1. 创建基础表结构
        
        # roles 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS roles (
                id INT AUTO_INCREMENT PRIMARY KEY,
                role_name VARCHAR(50) NOT NULL,
                role_type VARCHAR(20) NOT NULL UNIQUE
            )
        """)
        
        # 插入默认角色
        cursor.execute("SELECT COUNT(*) FROM roles")
        if cursor.fetchone()[0] == 0:
            cursor.executemany("""
                INSERT INTO roles (role_name, role_type) VALUES (%s, %s)
            """, [
                ('超级管理员', 'super_admin'),
                ('管理员', 'admin'),
                ('HR', 'hr'),
                ('IT运维', 'it'),
                ('普通用户', 'user')
            ])
            
        # users 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INT AUTO_INCREMENT PRIMARY KEY,
                username VARCHAR(50) NOT NULL UNIQUE,
                password VARCHAR(100) NOT NULL,
                real_name VARCHAR(50) NOT NULL,
                email VARCHAR(100),
                phone VARCHAR(20),
                role_id INT NOT NULL,
                department VARCHAR(50),
                status VARCHAR(20) DEFAULT 'active',
                current_session_id VARCHAR(64) DEFAULT NULL,
                created_by INT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (role_id) REFERENCES roles(id)
            )
        """)
        
        # 插入默认管理员
        cursor.execute("SELECT COUNT(*) FROM users")
        if cursor.fetchone()[0] == 0:
            cursor.execute("SELECT id FROM roles WHERE role_type = 'super_admin'")
            role_id = cursor.fetchone()[0]
            cursor.execute("""
                INSERT INTO users (username, password, real_name, role_id) 
                VALUES ('admin', 'admin123', '超级管理员', %s)
            """, (role_id,))

        # ticket_categories 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS ticket_categories (
                id INT AUTO_INCREMENT PRIMARY KEY,
                category_name VARCHAR(50) NOT NULL,
                form_template TEXT DEFAULT NULL COMMENT '自定义表单模板(JSON)',
                workflow_template TEXT DEFAULT NULL COMMENT '审批流程节点配置(JSON)'
            )
        """)
        
        # 插入默认分类
        cursor.execute("SELECT COUNT(*) FROM ticket_categories")
        if cursor.fetchone()[0] == 0:
            cursor.executemany("INSERT INTO ticket_categories (category_name) VALUES (%s)", [
                ('硬件故障',), ('软件故障',), ('网络故障',), ('打印机故障',), ('其他',)
            ])

        # tickets 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS tickets (
                id INT AUTO_INCREMENT PRIMARY KEY,
                ticket_no VARCHAR(50) NOT NULL UNIQUE,
                title VARCHAR(100) NOT NULL,
                category_id INT NOT NULL,
                description TEXT,
                equipment_type VARCHAR(50),
                location VARCHAR(100),
                submitter_id INT NOT NULL,
                submitter_name VARCHAR(50),
                priority VARCHAR(20) DEFAULT 'medium',
                extra_data TEXT DEFAULT NULL COMMENT '自定义表单数据(JSON)',
                workflow_snapshot TEXT DEFAULT NULL COMMENT '创建工单时的审批流程快照(JSON)',
                status VARCHAR(20) DEFAULT 'pending',
                assigned_to INT,
                claimed_at TIMESTAMP NULL,
                solution TEXT,
                completed_at TIMESTAMP NULL,
                closed_at TIMESTAMP NULL,
                satisfaction INT DEFAULT NULL COMMENT '满意度评分(1-5)',
                satisfaction_comment TEXT DEFAULT NULL COMMENT '满意度评价',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (category_id) REFERENCES ticket_categories(id),
                FOREIGN KEY (submitter_id) REFERENCES users(id),
                FOREIGN KEY (assigned_to) REFERENCES users(id)
            )
        """)

        # ticket_logs 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS ticket_logs (
                id INT AUTO_INCREMENT PRIMARY KEY,
                ticket_id INT NOT NULL,
                user_id INT NOT NULL,
                user_name VARCHAR(50),
                action VARCHAR(50) NOT NULL,
                content TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (ticket_id) REFERENCES tickets(id),
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        # login_logs 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS login_logs (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT,
                username VARCHAR(50),
                real_name VARCHAR(50),
                status VARCHAR(20) NOT NULL,
                ip_address VARCHAR(100),
                user_agent TEXT,
                message TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # operation_logs 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS operation_logs (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT,
                username VARCHAR(50),
                real_name VARCHAR(50),
                module VARCHAR(50),
                action VARCHAR(100) NOT NULL,
                detail TEXT,
                ip_address VARCHAR(100),
                user_agent TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # notifications 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS notifications (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL,
                ticket_id INT,
                title VARCHAR(100) NOT NULL,
                content TEXT,
                type VARCHAR(50),
                is_read TINYINT DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        # category_it_mapping 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS category_it_mapping (
                id INT AUTO_INCREMENT PRIMARY KEY,
                category_id INT NOT NULL,
                it_user_id INT NOT NULL,
                priority INT DEFAULT 0,
                is_active TINYINT DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (category_id) REFERENCES ticket_categories(id),
                FOREIGN KEY (it_user_id) REFERENCES users(id)
            )
        """)
        
        # webhook_configs 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS webhook_configs (
                id INT AUTO_INCREMENT PRIMARY KEY,
                user_id INT NOT NULL COMMENT '所属用户ID',
                name VARCHAR(100) NOT NULL COMMENT '配置名称',
                platform VARCHAR(20) NOT NULL COMMENT '平台类型：dingtalk/feishu/wechat',
                webhook_url TEXT NOT NULL COMMENT 'Webhook地址',
                secret VARCHAR(200) DEFAULT NULL COMMENT '签名密钥',
                enabled TINYINT DEFAULT 1 COMMENT '是否启用',
                scope VARCHAR(20) DEFAULT 'personal' COMMENT '配置范围：personal/global',
                notify_events JSON DEFAULT NULL COMMENT '通知事件类型',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id)
            )
        """)

        # system_settings 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS system_settings (
                setting_key VARCHAR(100) PRIMARY KEY,
                setting_value TEXT,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
            )
        """)

        default_settings = [
            ('ticket_no_prefix', 'TK'),
            ('ticket_no_date_format', '%Y%m%d%H%M%S'),
            ('ticket_no_random_digits', '3')
        ]
        for key, value in default_settings:
            cursor.execute("""
                INSERT INTO system_settings (setting_key, setting_value)
                VALUES (%s, %s)
                ON DUPLICATE KEY UPDATE setting_key = setting_key
            """, (key, value))

        # webhook_templates 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS webhook_templates (
                id INT AUTO_INCREMENT PRIMARY KEY,
                name VARCHAR(100) NOT NULL,
                event_type VARCHAR(50) NOT NULL,
                title_template VARCHAR(200) NOT NULL,
                content_template TEXT NOT NULL,
                is_active TINYINT DEFAULT 0,
                created_by INT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
                FOREIGN KEY (created_by) REFERENCES users(id)
            )
        """)

        cursor.execute("SELECT COUNT(*) FROM webhook_templates")
        if cursor.fetchone()[0] == 0:
            rows = [
                (
                    tpl['name'],
                    event_type,
                    tpl['title'],
                    tpl['content'],
                    1,
                    None
                )
                for event_type, tpl in DEFAULT_WEBHOOK_TEMPLATES.items()
            ]
            cursor.executemany("""
                INSERT INTO webhook_templates
                    (name, event_type, title_template, content_template, is_active, created_by)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, rows)
        
        # 检查并删除旧的created_by字段（用于升级旧表）
        cursor.execute("SHOW COLUMNS FROM webhook_configs LIKE 'created_by'")
        if cursor.fetchone():
            print("Removing old created_by column from webhook_configs...")
            # 先删除外键约束
            try:
                cursor.execute("ALTER TABLE webhook_configs DROP FOREIGN KEY webhook_configs_ibfk_1")
            except:
                pass
            # 再删除字段
            try:
                cursor.execute("ALTER TABLE webhook_configs DROP COLUMN created_by")
            except:
                pass
        
        # 2. 检查并添加字段（用于升级旧表）

        # 检查 current_session_id 字段
        cursor.execute("SHOW COLUMNS FROM users LIKE 'current_session_id'")
        if not cursor.fetchone():
            print("Adding current_session_id column to users...")
            cursor.execute("ALTER TABLE users ADD COLUMN current_session_id VARCHAR(64) DEFAULT NULL")
        
        # 检查 satisfaction 字段
        cursor.execute("SHOW COLUMNS FROM tickets LIKE 'satisfaction'")
        if not cursor.fetchone():
            print("Adding satisfaction column...")
            cursor.execute("ALTER TABLE tickets ADD COLUMN satisfaction INT DEFAULT NULL COMMENT '满意度评分(1-5)'")
            cursor.execute("ALTER TABLE tickets ADD COLUMN satisfaction_comment TEXT DEFAULT NULL COMMENT '满意度评价'")
            
        # 检查 form_template 字段
        cursor.execute("SHOW COLUMNS FROM ticket_categories LIKE 'form_template'")
        if not cursor.fetchone():
            print("Adding form_template column to ticket_categories...")
            cursor.execute("ALTER TABLE ticket_categories ADD COLUMN form_template TEXT DEFAULT NULL COMMENT '自定义表单模板(JSON)'")

        # 检查 extra_data 字段
        cursor.execute("SHOW COLUMNS FROM tickets LIKE 'extra_data'")
        if not cursor.fetchone():
            print("Adding extra_data column to tickets...")
            cursor.execute("ALTER TABLE tickets ADD COLUMN extra_data TEXT DEFAULT NULL COMMENT '自定义表单数据(JSON)'")

        # 检查 workflow_template 字段
        cursor.execute("SHOW COLUMNS FROM ticket_categories LIKE 'workflow_template'")
        if not cursor.fetchone():
            print("Adding workflow_template column to ticket_categories...")
            cursor.execute("ALTER TABLE ticket_categories ADD COLUMN workflow_template TEXT DEFAULT NULL COMMENT '审批流程节点配置(JSON)'")

        # 检查 workflow_snapshot 字段
        cursor.execute("SHOW COLUMNS FROM tickets LIKE 'workflow_snapshot'")
        if not cursor.fetchone():
            print("Adding workflow_snapshot column to tickets...")
            cursor.execute("ALTER TABLE tickets ADD COLUMN workflow_snapshot TEXT DEFAULT NULL COMMENT '创建工单时的审批流程快照(JSON)'")

        # 检查 webhook scope 字段
        cursor.execute("SHOW COLUMNS FROM webhook_configs LIKE 'scope'")
        if not cursor.fetchone():
            print("Adding scope column to webhook_configs...")
            cursor.execute("ALTER TABLE webhook_configs ADD COLUMN scope VARCHAR(20) DEFAULT 'personal' COMMENT '配置范围：personal/global'")
            
        conn.commit()
    except Exception as e:
        print(f"Migration failed: {e}")
    finally:
        conn.close()

# 执行迁移
check_and_migrate_db()

# Pydantic模型
class UserLogin(BaseModel):
    username: str
    password: str

class UserCreate(BaseModel):
    username: str
    password: str
    real_name: str
    email: Optional[str] = None
    phone: Optional[str] = None
    role_id: int
    department: Optional[str] = None

class TicketCreate(BaseModel):
    title: str
    category_id: int
    description: str
    equipment_type: str
    location: Optional[str] = None
    priority: str = "medium"
    extra_data: Optional[str] = None

class CategoryUpdate(BaseModel):
    category_name: Optional[str] = None
    form_template: Optional[str] = None
    workflow_template: Optional[str] = None

class CategoryCreate(BaseModel):
    category_name: str
    form_template: Optional[str] = None
    workflow_template: Optional[str] = None

class CategoryWorkflowUpdate(BaseModel):
    workflow_template: Optional[str] = None

class WorkflowAction(BaseModel):
    action: str
    comment: Optional[str] = None

class TicketRate(BaseModel):
    satisfaction: int
    comment: Optional[str] = None

class TicketUpdate(BaseModel):
    status: Optional[str] = None
    assigned_to: Optional[int] = None
    solution: Optional[str] = None

class CategoryITMapping(BaseModel):
    category_id: int
    it_user_id: int
    priority: int = 0

class WebhookConfigCreate(BaseModel):
    name: str
    platform: str  # dingtalk, feishu, wechat
    webhook_url: str
    secret: Optional[str] = None
    scope: Optional[str] = "personal"
    notify_events: Optional[List[str]] = None

class WebhookConfigUpdate(BaseModel):
    name: Optional[str] = None
    webhook_url: Optional[str] = None
    secret: Optional[str] = None
    enabled: Optional[int] = None
    scope: Optional[str] = None
    notify_events: Optional[List[str]] = None

class PasswordChange(BaseModel):
    old_password: str
    new_password: str

class TicketNumberSettings(BaseModel):
    prefix: str = "TK"
    date_format: str = "%Y%m%d%H%M%S"
    random_digits: int = 3

class WebhookTemplateCreate(BaseModel):
    name: str
    event_type: str
    title_template: str
    content_template: str
    is_active: int = 0

class WebhookTemplateUpdate(BaseModel):
    name: Optional[str] = None
    event_type: Optional[str] = None
    title_template: Optional[str] = None
    content_template: Optional[str] = None
    is_active: Optional[int] = None

# JWT工具函数
def create_access_token(data: dict):
    to_encode = data.copy()
    expire = datetime.utcnow() + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

def verify_token(token: str):
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        session_id = payload.get("session_id")
        if not session_id:
            return None

        conn = pymysql.connect(**DB_CONFIG)
        try:
            if DB_TIME_ZONE:
                with conn.cursor() as cursor:
                    cursor.execute("SET time_zone = %s", (DB_TIME_ZONE,))
            cursor = conn.cursor(pymysql.cursors.DictCursor)
            cursor.execute("SELECT current_session_id, status FROM users WHERE id = %s", (payload.get("user_id"),))
            user = cursor.fetchone()
            if not user or user['status'] != 'active' or user['current_session_id'] != session_id:
                return None
        finally:
            conn.close()
        return payload
    except:
        return None

def get_request_ip(request: Optional[Request]) -> Optional[str]:
    if not request:
        return None
    forwarded_for = request.headers.get("x-forwarded-for")
    if forwarded_for:
        return forwarded_for.split(",")[0].strip()
    return request.client.host if request.client else None

def get_request_user_agent(request: Optional[Request]) -> Optional[str]:
    return request.headers.get("user-agent") if request else None

def log_login(conn, user: Optional[dict], username: str, status: str, message: str, request: Optional[Request] = None):
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO login_logs (user_id, username, real_name, status, ip_address, user_agent, message)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
    """, (
        user['id'] if user else None,
        username,
        user.get('real_name') if user else None,
        status,
        get_request_ip(request),
        get_request_user_agent(request),
        message
    ))

def log_operation(conn, user_id: Optional[int], username: Optional[str], real_name: Optional[str], module: str, action: str, detail: str, request: Optional[Request] = None):
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO operation_logs (user_id, username, real_name, module, action, detail, ip_address, user_agent)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
    """, (
        user_id,
        username,
        real_name,
        module,
        action,
        detail,
        get_request_ip(request),
        get_request_user_agent(request)
    ))

def get_db_now(cursor):
    cursor.execute("SELECT NOW() AS now_time")
    row = cursor.fetchone()
    if isinstance(row, dict):
        return row['now_time']
    return row[0]

def now_text(conn=None) -> str:
    if conn is not None:
        cursor = conn.cursor(pymysql.cursors.DictCursor)
        return get_db_now(cursor).strftime('%Y-%m-%d %H:%M:%S')
    return datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S')

def render_text_template(template: str, variables: dict) -> str:
    class SafeDict(dict):
        def __missing__(self, key):
            return "{" + key + "}"

    return template.format_map(SafeDict({k: "" if v is None else v for k, v in variables.items()}))

def get_setting(cursor, key: str, default: str) -> str:
    cursor.execute("SELECT setting_value FROM system_settings WHERE setting_key = %s", (key,))
    row = cursor.fetchone()
    if isinstance(row, dict):
        return row['setting_value'] or default
    if row:
        return row[0] or default
    return default

def generate_ticket_no(cursor) -> str:
    prefix = get_setting(cursor, 'ticket_no_prefix', 'TK')
    date_format = get_setting(cursor, 'ticket_no_date_format', '%Y%m%d%H%M%S')
    random_digits_raw = get_setting(cursor, 'ticket_no_random_digits', '3')
    try:
        random_digits = max(1, min(12, int(random_digits_raw)))
    except ValueError:
        random_digits = 3

    db_now = get_db_now(cursor)

    for _ in range(10):
        random_suffix = ''.join(random.choices(string.digits, k=random_digits))
        ticket_no = f"{prefix}{db_now.strftime(date_format)}{random_suffix}"
        cursor.execute("SELECT id FROM tickets WHERE ticket_no = %s", (ticket_no,))
        if not cursor.fetchone():
            return ticket_no

    random_suffix = ''.join(random.choices(string.digits, k=random_digits + 3))
    return f"{prefix}{db_now.strftime(date_format)}{random_suffix}"

def require_admin(payload: dict, conn):
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    if not user or user['role_type'] not in ['super_admin', 'admin']:
        raise HTTPException(status_code=403, detail="只有管理员可以操作")
    return user

def normalize_id_list(value: Any) -> List[int]:
    if not isinstance(value, list):
        return []

    ids = []
    seen = set()
    for item in value:
        try:
            user_id = int(item)
        except (TypeError, ValueError):
            continue
        if user_id > 0 and user_id not in seen:
            seen.add(user_id)
            ids.append(user_id)
    return ids

def validate_workflow_template(template_value: Optional[str], conn) -> Optional[str]:
    if not template_value:
        return None

    try:
        template = json.loads(template_value) if isinstance(template_value, str) else template_value
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="流程配置不是有效的JSON")

    nodes = template.get("nodes") if isinstance(template, dict) else None
    if not isinstance(nodes, list):
        raise HTTPException(status_code=400, detail="流程配置必须包含nodes数组")

    normalized_nodes = []
    referenced_user_ids = set()
    for index, node in enumerate(nodes, start=1):
        if not isinstance(node, dict):
            raise HTTPException(status_code=400, detail=f"第{index}个节点格式不正确")

        name = str(node.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=400, detail=f"第{index}个节点名称不能为空")

        approver_ids = normalize_id_list(node.get("approver_ids"))
        cc_ids = normalize_id_list(node.get("cc_ids"))
        if not approver_ids:
            raise HTTPException(status_code=400, detail=f"节点“{name}”至少需要一个审批人")

        referenced_user_ids.update(approver_ids)
        referenced_user_ids.update(cc_ids)
        normalized_nodes.append({
            "id": str(node.get("id") or f"node_{index}"),
            "name": name[:50],
            "approver_ids": approver_ids,
            "cc_ids": cc_ids,
            "comment": str(node.get("comment") or "").strip()[:500],
            "order": index
        })

    if referenced_user_ids:
        cursor = conn.cursor(pymysql.cursors.DictCursor)
        placeholders = ",".join(["%s"] * len(referenced_user_ids))
        cursor.execute(f"SELECT id FROM users WHERE status = 'active' AND id IN ({placeholders})", tuple(referenced_user_ids))
        active_user_ids = {row['id'] for row in cursor.fetchall()}
        invalid_user_ids = sorted(referenced_user_ids - active_user_ids)
        if invalid_user_ids:
            raise HTTPException(status_code=400, detail=f"流程中包含无效或停用用户：{', '.join(map(str, invalid_user_ids))}")

    return json.dumps({"nodes": normalized_nodes}, ensure_ascii=False)

def load_workflow_people(cursor, workflow_template: Optional[str]) -> Dict[str, Any]:
    if not workflow_template:
        return {"nodes": [], "user_ids": []}

    try:
        workflow = json.loads(workflow_template)
    except (TypeError, json.JSONDecodeError):
        return {"nodes": [], "user_ids": []}

    nodes = workflow.get("nodes") if isinstance(workflow, dict) else []
    if not isinstance(nodes, list):
        return {"nodes": [], "user_ids": []}

    all_user_ids = []
    seen = set()
    for node in nodes:
        if not isinstance(node, dict):
            continue
        for user_id in normalize_id_list(node.get("approver_ids")) + normalize_id_list(node.get("cc_ids")):
            if user_id not in seen:
                seen.add(user_id)
                all_user_ids.append(user_id)

    user_map = {}
    if all_user_ids:
        placeholders = ",".join(["%s"] * len(all_user_ids))
        cursor.execute(f"SELECT id, real_name, username FROM users WHERE id IN ({placeholders})", tuple(all_user_ids))
        user_map = {row['id']: row for row in cursor.fetchall()}

    hydrated_nodes = []
    for node in nodes:
        if not isinstance(node, dict):
            continue
        approver_ids = normalize_id_list(node.get("approver_ids"))
        cc_ids = normalize_id_list(node.get("cc_ids"))
        hydrated_nodes.append({
            **node,
            "approvers": [user_map[user_id] for user_id in approver_ids if user_id in user_map],
            "cc_users": [user_map[user_id] for user_id in cc_ids if user_id in user_map]
        })

    return {"nodes": hydrated_nodes, "user_ids": all_user_ids}

def build_workflow_snapshot(nodes: List[Dict[str, Any]]) -> Optional[str]:
    if not nodes:
        return None

    snapshot_nodes = []
    for index, node in enumerate(nodes):
        snapshot_nodes.append({
            **node,
            "status": "pending" if index == 0 else "waiting",
            "approved_by": None,
            "approved_by_name": None,
            "approved_at": None,
            "approval_comment": None
        })
    return json.dumps({"current_index": 0, "status": "pending", "nodes": snapshot_nodes}, ensure_ascii=False)

def parse_workflow_snapshot(snapshot: Optional[str]) -> Optional[Dict[str, Any]]:
    if not snapshot:
        return None
    try:
        workflow = json.loads(snapshot)
    except (TypeError, json.JSONDecodeError):
        return None
    if not isinstance(workflow, dict) or not isinstance(workflow.get("nodes"), list):
        return None

    nodes = workflow["nodes"]
    current_index = workflow.get("current_index")
    if current_index is None:
        current_index = 0
        for idx, node in enumerate(nodes):
            if node.get("status") in ["pending", None]:
                current_index = idx
                break
    workflow["current_index"] = current_index
    workflow["status"] = workflow.get("status") or "pending"
    return workflow

def workflow_user_ids(workflow: Optional[Dict[str, Any]]) -> set:
    if not workflow:
        return set()
    user_ids = set()
    for node in workflow.get("nodes", []):
        if not isinstance(node, dict):
            continue
        user_ids.update(normalize_id_list(node.get("approver_ids")))
        user_ids.update(normalize_id_list(node.get("cc_ids")))
    return user_ids

def apply_ticket_workflow_fields(ticket: Dict[str, Any], user_id: int) -> Dict[str, Any]:
    workflow = parse_workflow_snapshot(ticket.get('workflow_snapshot'))
    ticket['workflow'] = workflow
    ticket['can_approve_workflow'] = False
    ticket['current_workflow_node'] = None
    ticket['workflow_current_node_name'] = None
    ticket['workflow_relation'] = None

    if not workflow:
        return ticket

    if user_id in workflow_user_ids(workflow):
        ticket['workflow_relation'] = 'participant'

    current_index = workflow.get('current_index', 0)
    nodes = workflow.get('nodes', [])
    if 0 <= current_index < len(nodes):
        current_node = nodes[current_index]
        ticket['current_workflow_node'] = current_node
        ticket['workflow_current_node_name'] = current_node.get('name')
        approver_ids = normalize_id_list(current_node.get('approver_ids'))
        cc_ids = normalize_id_list(current_node.get('cc_ids'))
        if ticket.get('status') == 'pending_approval' and user_id in approver_ids:
            ticket['can_approve_workflow'] = True
            ticket['workflow_relation'] = 'current_approver'
        elif user_id in cc_ids:
            ticket['workflow_relation'] = ticket['workflow_relation'] or 'cc'

    return ticket

def can_user_see_ticket(ticket: Dict[str, Any], user_id: int) -> bool:
    if ticket.get('submitter_id') == user_id or ticket.get('assigned_to') == user_id:
        return True
    workflow = parse_workflow_snapshot(ticket.get('workflow_snapshot'))
    return user_id in workflow_user_ids(workflow)

def summarize_workflow_nodes(nodes: List[Dict[str, Any]]) -> str:
    node_lines = []
    for node in nodes:
        approver_names = "、".join([approver['real_name'] for approver in node.get("approvers", [])]) or "未配置"
        cc_names = "、".join([cc_user['real_name'] for cc_user in node.get("cc_users", [])]) or "无"
        comment = f"，说明：{node.get('comment')}" if node.get('comment') else ""
        node_lines.append(f"{node.get('order')}. {node.get('name')}：审批人 {approver_names}，抄送 {cc_names}{comment}")
    return "\n".join(node_lines)

async def assign_ticket_after_workflow(cursor, ticket_id: int, ticket: Dict[str, Any], submitter_name: str):
    cursor.execute("""
        SELECT m.it_user_id, u.real_name as it_user_name
        FROM category_it_mapping m
        JOIN users u ON m.it_user_id = u.id
        WHERE m.category_id = %s AND m.is_active = 1
        ORDER BY m.priority DESC
        LIMIT 1
    """, (ticket['category_id'],))
    assigned_it = cursor.fetchone()

    if assigned_it:
        cursor.execute("""
            UPDATE tickets 
            SET status = 'claimed', assigned_to = %s, claimed_at = NOW()
            WHERE id = %s
        """, (assigned_it['it_user_id'], ticket_id))
        cursor.execute("""
            INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
            VALUES (%s, %s, %s, %s, %s)
        """, (ticket_id, assigned_it['it_user_id'], assigned_it['it_user_name'], '自动分配', 
              f"审批通过后，系统自动分配给{assigned_it['it_user_name']}"))
        cursor.execute("""
            INSERT INTO notifications (user_id, ticket_id, title, content, type)
            VALUES (%s, %s, %s, %s, %s)
        """, (assigned_it['it_user_id'], ticket_id, '新工单分配', 
              f"{submitter_name}提交的工单：{ticket['title']} 已审批通过并分配给您", 'new_ticket'))
        await manager.send_personal_message({
            "type": "new_ticket_assigned",
            "ticket_id": ticket_id,
            "ticket_no": ticket['ticket_no'],
            "title": ticket['title'],
            "message": f"工单已审批通过并自动分配给您"
        }, assigned_it['it_user_id'])
        return assigned_it

    cursor.execute("UPDATE tickets SET status = 'pending' WHERE id = %s", (ticket_id,))
    cursor.execute("SELECT u.id FROM users u JOIN roles r ON u.role_id = r.id WHERE r.role_type = 'it'")
    it_users = cursor.fetchall()
    for it_user in it_users:
        cursor.execute("""
            INSERT INTO notifications (user_id, ticket_id, title, content, type)
            VALUES (%s, %s, %s, %s, %s)
        """, (it_user['id'], ticket_id, '新工单提醒', f"{submitter_name}提交的工单：{ticket['title']} 已审批通过，请处理", 'new_ticket'))
        await manager.send_personal_message({
            "type": "new_ticket",
            "ticket_id": ticket_id,
            "ticket_no": ticket['ticket_no'],
            "title": ticket['title'],
            "message": "工单已审批通过，请处理"
        }, it_user['id'])
    return None

# 根路径重定向
@app.get("/")
async def root():
    return RedirectResponse(url="/static/index.html")

# 登录接口
@app.post("/api/login")
async def login(user: UserLogin, request: Request, conn=Depends(get_db)):
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT u.*, r.role_type, r.role_name 
        FROM users u 
        JOIN roles r ON u.role_id = r.id 
        WHERE u.username = %s AND u.status = 'active'
    """, (user.username,))
    db_user = cursor.fetchone()
    
    if not db_user:
        log_login(conn, None, user.username, 'failed', '用户名不存在或账号未启用', request)
        conn.commit()
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    
    # 验证密码 - 简化为明文比较（临时方案，确保功能可用）
    if user.password != db_user['password']:
        log_login(conn, db_user, user.username, 'failed', '密码错误', request)
        conn.commit()
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    
    session_id = uuid.uuid4().hex
    cursor.execute("UPDATE users SET current_session_id = %s WHERE id = %s", (session_id, db_user['id']))
    log_login(conn, db_user, user.username, 'success', '登录成功，旧会话已失效', request)
    log_operation(conn, db_user['id'], db_user['username'], db_user['real_name'], 'auth', '登录', '用户登录系统', request)
    conn.commit()

    token = create_access_token({"user_id": db_user['id'], "username": db_user['username'], "session_id": session_id})
    
    return {
        "token": token,
        "user": {
            "id": db_user['id'],
            "username": db_user['username'],
            "real_name": db_user['real_name'],
            "email": db_user['email'],
            "role_id": db_user['role_id'],
            "role_type": db_user['role_type'],
            "role_name": db_user['role_name'],
            "department": db_user['department']
        }
    }

# 获取当前用户
@app.get("/api/current-user")
async def get_current_user(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT u.*, r.role_type, r.role_name 
        FROM users u 
        JOIN roles r ON u.role_id = r.id 
        WHERE u.id = %s
    """, (payload['user_id'],))
    user = cursor.fetchone()
    
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    
    return {
        "id": user['id'],
        "username": user['username'],
        "real_name": user['real_name'],
        "email": user['email'],
        "role_id": user['role_id'],
        "role_type": user['role_type'],
        "role_name": user['role_name'],
        "department": user['department']
    }

@app.post("/api/logout")
async def logout(token: str, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        return {"message": "已退出"}

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    if user:
        cursor.execute("UPDATE users SET current_session_id = NULL WHERE id = %s", (payload['user_id'],))
        log_operation(conn, user['id'], user['username'], user['real_name'], 'auth', '退出登录', '用户主动退出登录', request)
        conn.commit()

    return {"message": "已退出"}

@app.put("/api/current-user/password")
async def change_current_user_password(token: str, data: PasswordChange, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")

    if len(data.new_password.strip()) < 6:
        raise HTTPException(status_code=400, detail="新密码至少需要6位")

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM users WHERE id = %s AND status = 'active'", (payload['user_id'],))
    user = cursor.fetchone()

    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")

    if data.old_password != user['password']:
        raise HTTPException(status_code=400, detail="原密码不正确")

    cursor.execute("UPDATE users SET password = %s, current_session_id = NULL WHERE id = %s", (data.new_password, payload['user_id']))
    log_operation(conn, user['id'], user['username'], user['real_name'], 'auth', '修改密码', '用户修改密码成功，当前会话已强制失效', request)
    conn.commit()

    return {"message": "密码修改成功"}

# 用户管理接口
@app.get("/api/users")
async def get_users(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT u.*, r.role_name, r.role_type,
               creator.real_name as creator_name
        FROM users u
        JOIN roles r ON u.role_id = r.id
        LEFT JOIN users creator ON u.created_by = creator.id
        ORDER BY u.created_at DESC
    """)
    users = cursor.fetchall()
    return {"users": users}

@app.post("/api/users")
async def create_user(token: str, user: UserCreate, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 检查创建者权限
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    creator = cursor.fetchone()
    
    if creator['role_type'] not in ['super_admin', 'admin', 'hr', 'it']:
        raise HTTPException(status_code=403, detail="没有权限创建用户")
    
    # HR只能创建普通用户
    if creator['role_type'] == 'hr':
        cursor.execute("SELECT role_type FROM roles WHERE id = %s", (user.role_id,))
        target_role = cursor.fetchone()
        if not target_role or target_role['role_type'] != 'user':
            raise HTTPException(status_code=403, detail="HR只能创建普通用户")
    
    # 检查用户名是否存在
    cursor.execute("SELECT id FROM users WHERE username = %s", (user.username,))
    if cursor.fetchone():
        raise HTTPException(status_code=400, detail="用户名已存在")
    
    # 密码直接使用明文（临时方案）
    password = user.password
    
    # 创建用户
    cursor.execute("""
        INSERT INTO users (username, password, real_name, email, phone, role_id, department, created_by)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
    """, (user.username, password, user.real_name, user.email, user.phone, user.role_id, user.department, payload['user_id']))
    new_user_id = cursor.lastrowid
    log_operation(conn, creator['id'], creator['username'], creator['real_name'], 'user', '创建用户', f"创建用户：{user.username} / {user.real_name}", request)
    
    conn.commit()
    
    return {"message": "用户创建成功", "user_id": new_user_id}

@app.post("/api/users/import")
async def import_users_from_excel(token: str, file: UploadFile = File(...), conn=Depends(get_db)):
    """批量导入用户（Excel文件）"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 检查创建者权限
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    creator = cursor.fetchone()
    
    if creator['role_type'] not in ['super_admin', 'admin', 'hr', 'it']:
        raise HTTPException(status_code=403, detail="没有权限导入用户")
    
    # 读取Excel文件
    try:
        contents = await file.read()
        workbook = openpyxl.load_workbook(BytesIO(contents))
        sheet = workbook.active
        
        success_count = 0
        error_list = []
        
        # 跳过表头，从第二行开始读取
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:  # 如果用户名为空，跳过
                continue
            
            try:
                username = str(row[0]).strip()
                password = str(row[1]).strip() if row[1] else "123456"  # 默认密码
                real_name = str(row[2]).strip() if row[2] else username
                email = str(row[3]).strip() if row[3] else None
                phone = str(row[4]).strip() if row[4] else None
                role_name = str(row[5]).strip() if row[5] else "普通用户"
                department = str(row[6]).strip() if row[6] else None
                
                # 根据角色名称查找role_id
                cursor.execute("SELECT id FROM roles WHERE role_name = %s", (role_name,))
                role = cursor.fetchone()
                if not role:
                    error_list.append(f"第{row_idx}行：角色'{role_name}'不存在")
                    continue
                
                role_id = role['id']
                
                # 检查用户名是否已存在
                cursor.execute("SELECT id FROM users WHERE username = %s", (username,))
                if cursor.fetchone():
                    error_list.append(f"第{row_idx}行：用户名'{username}'已存在")
                    continue
                
                # 创建用户
                cursor.execute("""
                    INSERT INTO users (username, password, real_name, email, phone, role_id, department, created_by)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                """, (username, password, real_name, email, phone, role_id, department, payload['user_id']))
                
                success_count += 1
                
            except Exception as e:
                error_list.append(f"第{row_idx}行：{str(e)}")
                continue
        
        conn.commit()
        
        return {
            "message": f"导入完成，成功{success_count}条",
            "success_count": success_count,
            "error_count": len(error_list),
            "errors": error_list
        }
        
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{str(e)}")

# 角色管理接口
@app.get("/api/roles")
async def get_roles(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM roles ORDER BY id")
    roles = cursor.fetchall()
    return {"roles": roles}

# 工单分类与IT人员配置接口
@app.get("/api/category-it-mapping")
async def get_category_it_mapping(token: str, conn=Depends(get_db)):
    """获取工单分类与IT人员映射配置"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT m.*, c.category_name, u.real_name as it_user_name, u.username as it_username
        FROM category_it_mapping m
        JOIN ticket_categories c ON m.category_id = c.id
        JOIN users u ON m.it_user_id = u.id
        WHERE m.is_active = 1
        ORDER BY m.category_id, m.priority DESC
    """)
    mappings = cursor.fetchall()
    return {"mappings": mappings}

@app.post("/api/category-it-mapping")
async def create_category_it_mapping(token: str, mapping: CategoryITMapping, conn=Depends(get_db)):
    """创建工单分类与IT人员映射"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 检查权限
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    if user['role_type'] not in ['super_admin', 'admin']:
        raise HTTPException(status_code=403, detail="没有权限配置")
    
    # 创建映射
    cursor.execute("""
        INSERT INTO category_it_mapping (category_id, it_user_id, priority)
        VALUES (%s, %s, %s)
    """, (mapping.category_id, mapping.it_user_id, mapping.priority))
    
    conn.commit()
    
    return {"message": "配置成功", "mapping_id": cursor.lastrowid}

@app.delete("/api/category-it-mapping/{mapping_id}")
async def delete_category_it_mapping(mapping_id: int, token: str, conn=Depends(get_db)):
    """删除工单分类与IT人员映射"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 检查权限
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    if user['role_type'] not in ['super_admin', 'admin']:
        raise HTTPException(status_code=403, detail="没有权限删除配置")
    
    # 删除映射
    cursor.execute("UPDATE category_it_mapping SET is_active = 0 WHERE id = %s", (mapping_id,))
    
    conn.commit()
    
    return {"message": "删除成功"}

# 工单分类接口
@app.get("/api/categories")
async def get_categories(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM ticket_categories ORDER BY id")
    categories = cursor.fetchall()
    return {"categories": categories}

@app.post("/api/categories")
async def create_category(token: str, category: CategoryCreate, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    user = require_admin(payload, conn)
    category_name = category.category_name.strip()
    if not category_name:
        raise HTTPException(status_code=400, detail="工单类型名称不能为空")

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT id FROM ticket_categories WHERE category_name = %s", (category_name,))
    if cursor.fetchone():
        raise HTTPException(status_code=400, detail="工单类型已存在")

    workflow_template = validate_workflow_template(category.workflow_template, conn)
    cursor.execute("""
        INSERT INTO ticket_categories (category_name, form_template, workflow_template)
        VALUES (%s, %s, %s)
    """, (category_name, category.form_template, workflow_template))
    category_id = cursor.lastrowid
    log_operation(conn, user['id'], user['username'], user['real_name'], 'category', '创建工单类型', f"创建工单类型：{category_name}", request)
    conn.commit()

    return {"message": "工单类型创建成功", "category_id": category_id}

@app.put("/api/categories/{category_id}")
async def update_category(category_id: int, token: str, category: CategoryUpdate, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    user = require_admin(payload, conn)
    cursor = conn.cursor(pymysql.cursors.DictCursor)

    cursor.execute("SELECT * FROM ticket_categories WHERE id = %s", (category_id,))
    existing = cursor.fetchone()
    if not existing:
        raise HTTPException(status_code=404, detail="工单类型不存在")

    updates = []
    params = []
    if category.category_name is not None:
        category_name = category.category_name.strip()
        if not category_name:
            raise HTTPException(status_code=400, detail="工单类型名称不能为空")
        cursor.execute("SELECT id FROM ticket_categories WHERE category_name = %s AND id <> %s", (category_name, category_id))
        if cursor.fetchone():
            raise HTTPException(status_code=400, detail="工单类型已存在")
        updates.append("category_name = %s")
        params.append(category_name)

    if category.form_template is not None:
        updates.append("form_template = %s")
        params.append(category.form_template)

    if category.workflow_template is not None:
        updates.append("workflow_template = %s")
        params.append(validate_workflow_template(category.workflow_template, conn))

    if not updates:
        return {"message": "没有需要更新的内容"}

    params.append(category_id)
    cursor.execute(f"UPDATE ticket_categories SET {', '.join(updates)} WHERE id = %s", tuple(params))
    log_operation(conn, user['id'], user['username'], user['real_name'], 'category', '更新工单类型', f"更新工单类型：{existing['category_name']}", request)
    conn.commit()
    
    return {"message": "分类配置更新成功"}

@app.put("/api/categories/{category_id}/workflow")
async def update_category_workflow(category_id: int, token: str, workflow: CategoryWorkflowUpdate, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    user = require_admin(payload, conn)
    workflow_template = validate_workflow_template(workflow.workflow_template, conn)
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT category_name FROM ticket_categories WHERE id = %s", (category_id,))
    category = cursor.fetchone()
    if not category:
        raise HTTPException(status_code=404, detail="工单类型不存在")

    cursor.execute("UPDATE ticket_categories SET workflow_template = %s WHERE id = %s", (workflow_template, category_id))
    log_operation(conn, user['id'], user['username'], user['real_name'], 'workflow', '保存审批流程', f"保存“{category['category_name']}”审批流程", request)
    conn.commit()

    return {"message": "审批流程保存成功"}

@app.delete("/api/categories/{category_id}")
async def delete_category(category_id: int, token: str, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    user = require_admin(payload, conn)
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM ticket_categories WHERE id = %s", (category_id,))
    category = cursor.fetchone()
    if not category:
        raise HTTPException(status_code=404, detail="工单类型不存在")

    cursor.execute("SELECT COUNT(*) AS total FROM tickets WHERE category_id = %s", (category_id,))
    if cursor.fetchone()['total'] > 0:
        raise HTTPException(status_code=400, detail="该工单类型已有工单使用，不能删除")

    cursor.execute("DELETE FROM category_it_mapping WHERE category_id = %s", (category_id,))
    cursor.execute("DELETE FROM ticket_categories WHERE id = %s", (category_id,))
    log_operation(conn, user['id'], user['username'], user['real_name'], 'category', '删除工单类型', f"删除工单类型：{category['category_name']}", request)
    conn.commit()

    return {"message": "工单类型删除成功"}

# 工单管理接口
@app.get("/api/tickets")
async def get_tickets(token: str, status: Optional[str] = None, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取当前用户信息
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    current_user = cursor.fetchone()
    
    if current_user['role_type'] in ['super_admin', 'admin', 'hr', 'it']:
        if status:
            cursor.execute("""
                SELECT t.*, c.category_name, 
                       submitter.real_name as submitter_real_name,
                       assignee.real_name as assignee_name
                FROM tickets t
                JOIN ticket_categories c ON t.category_id = c.id
                JOIN users submitter ON t.submitter_id = submitter.id
                LEFT JOIN users assignee ON t.assigned_to = assignee.id
                WHERE t.status = %s
                ORDER BY t.created_at DESC
            """, (status,))
        else:
            cursor.execute("""
                SELECT t.*, c.category_name, 
                       submitter.real_name as submitter_real_name,
                       assignee.real_name as assignee_name
                FROM tickets t
                JOIN ticket_categories c ON t.category_id = c.id
                JOIN users submitter ON t.submitter_id = submitter.id
                LEFT JOIN users assignee ON t.assigned_to = assignee.id
                ORDER BY t.created_at DESC
            """)
        tickets = cursor.fetchall()
    else:
        # 普通用户可以看到自己提交的工单，以及自己参与审批/抄送的工单。
        if status:
            cursor.execute("""
                SELECT t.*, c.category_name, 
                       submitter.real_name as submitter_real_name,
                       assignee.real_name as assignee_name
                FROM tickets t
                JOIN ticket_categories c ON t.category_id = c.id
                JOIN users submitter ON t.submitter_id = submitter.id
                LEFT JOIN users assignee ON t.assigned_to = assignee.id
                WHERE (t.submitter_id = %s OR t.workflow_snapshot IS NOT NULL) AND t.status = %s
                ORDER BY t.created_at DESC
            """, (payload['user_id'], status))
        else:
            cursor.execute("""
                SELECT t.*, c.category_name, 
                       submitter.real_name as submitter_real_name,
                       assignee.real_name as assignee_name
                FROM tickets t
                JOIN ticket_categories c ON t.category_id = c.id
                JOIN users submitter ON t.submitter_id = submitter.id
                LEFT JOIN users assignee ON t.assigned_to = assignee.id
                WHERE t.submitter_id = %s OR t.workflow_snapshot IS NOT NULL
                ORDER BY t.created_at DESC
            """, (payload['user_id'],))
        tickets = [ticket for ticket in cursor.fetchall() if can_user_see_ticket(ticket, payload['user_id'])]
    
    tickets = [apply_ticket_workflow_fields(ticket, payload['user_id']) for ticket in tickets]
    return {"tickets": tickets}

@app.post("/api/tickets")
async def create_ticket(token: str, ticket: TicketCreate, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取用户信息
    cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
    user = cursor.fetchone()

    cursor.execute("SELECT category_name, workflow_template FROM ticket_categories WHERE id = %s", (ticket.category_id,))
    category = cursor.fetchone()
    if not category:
        raise HTTPException(status_code=400, detail="工单分类不存在")

    workflow_info = load_workflow_people(cursor, category.get('workflow_template'))
    workflow_snapshot = build_workflow_snapshot(workflow_info["nodes"])
    initial_status = "pending_approval" if workflow_snapshot else "pending"
    
    # 生成工单编号
    ticket_no = generate_ticket_no(cursor)
    
    # 创建工单
    cursor.execute("""
        INSERT INTO tickets (ticket_no, title, category_id, description, equipment_type, location, 
                           submitter_id, submitter_name, priority, extra_data, workflow_snapshot, status)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    """, (ticket_no, ticket.title, ticket.category_id, ticket.description, ticket.equipment_type, 
          ticket.location, payload['user_id'], user['real_name'], ticket.priority, ticket.extra_data, workflow_snapshot, initial_status))
    
    ticket_id = cursor.lastrowid
    
    # 记录日志
    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], '创建工单', f"创建了工单：{ticket.title}"))

    if workflow_info["nodes"]:
        cursor.execute("""
            INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
            VALUES (%s, %s, %s, %s, %s)
        """, (ticket_id, payload['user_id'], '系统', '发起审批', "已发起审批流程：\n" + summarize_workflow_nodes(workflow_info["nodes"])))

        workflow_data = parse_workflow_snapshot(workflow_snapshot)
        first_node = workflow_data["nodes"][0] if workflow_data and workflow_data["nodes"] else None
        first_approver_ids = normalize_id_list(first_node.get("approver_ids")) if first_node else []
        first_cc_ids = normalize_id_list(first_node.get("cc_ids")) if first_node else []
        for notify_user_id in first_approver_ids + first_cc_ids:
            if notify_user_id == payload['user_id']:
                continue
            notice_title = '待审批工单' if notify_user_id in first_approver_ids else '工单抄送'
            notice_content = f"{user['real_name']}提交了“{category['category_name']}”工单：{ticket.title}"
            if notify_user_id in first_approver_ids:
                notice_content += "，请您审批"
            else:
                notice_content += "，抄送给您"
            cursor.execute("""
                INSERT INTO notifications (user_id, ticket_id, title, content, type)
                VALUES (%s, %s, %s, %s, %s)
            """, (notify_user_id, ticket_id, notice_title, notice_content, 'workflow_notice'))
            await manager.send_personal_message({
                "type": "workflow_notice",
                "ticket_id": ticket_id,
                "ticket_no": ticket_no,
                "title": ticket.title,
                "message": notice_content
            }, notify_user_id)

        conn.commit()
        webhook_title, webhook_content = get_webhook_message('ticket_created', {
            "ticket_no": ticket_no,
            "title": ticket.title,
            "submitter": user['real_name'],
            "priority": ticket.priority,
            "time": now_text(conn)
        }, conn)
        await send_webhook_notification('ticket_created', webhook_title, webhook_content, [payload['user_id']] + first_approver_ids + first_cc_ids, conn)
        log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '创建工单', f"创建工单 {ticket_no}：{ticket.title}，进入审批", request)
        conn.commit()
        return {"message": "工单创建成功，已进入审批流程", "ticket_id": ticket_id, "ticket_no": ticket_no}

    # 根据工单分类查找对应的IT人员
    cursor.execute("""
        SELECT m.it_user_id, u.real_name as it_user_name
        FROM category_it_mapping m
        JOIN users u ON m.it_user_id = u.id
        WHERE m.category_id = %s AND m.is_active = 1
        ORDER BY m.priority DESC
        LIMIT 1
    """, (ticket.category_id,))
    
    assigned_it = cursor.fetchone()
    
    if assigned_it:
        # 如果找到对应的IT人员，自动分配
        cursor.execute("""
            UPDATE tickets 
            SET status = 'claimed', assigned_to = %s, claimed_at = NOW()
            WHERE id = %s
        """, (assigned_it['it_user_id'], ticket_id))
        
        # 记录自动分配日志
        cursor.execute("""
            INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
            VALUES (%s, %s, %s, %s, %s)
        """, (ticket_id, assigned_it['it_user_id'], assigned_it['it_user_name'], '自动分配', 
              f"系统自动分配给{assigned_it['it_user_name']}"))
        
        # 通知被分配的IT人员
        cursor.execute("""
            INSERT INTO notifications (user_id, ticket_id, title, content, type)
            VALUES (%s, %s, %s, %s, %s)
        """, (assigned_it['it_user_id'], ticket_id, '新工单分配', 
              f"{user['real_name']}提交了新工单：{ticket.title}，已自动分配给您", 'new_ticket'))
        
        # 发送WebSocket实时通知给IT人员
        await manager.send_personal_message({
            "type": "new_ticket_assigned",
            "ticket_id": ticket_id,
            "ticket_no": ticket_no,
            "title": ticket.title,
            "message": f"{user['real_name']}提交了新工单，已自动分配给您"
        }, assigned_it['it_user_id'])
        
    else:
        # 如果没有找到对应的IT人员，通知所有IT人员
        cursor.execute("SELECT u.id FROM users u JOIN roles r ON u.role_id = r.id WHERE r.role_type = 'it'")
        it_users = cursor.fetchall()
        
        for it_user in it_users:
            cursor.execute("""
                INSERT INTO notifications (user_id, ticket_id, title, content, type)
                VALUES (%s, %s, %s, %s, %s)
            """, (it_user['id'], ticket_id, '新工单提醒', f"{user['real_name']}提交了新工单：{ticket.title}", 'new_ticket'))
            
            # 发送WebSocket通知
            await manager.send_personal_message({
                "type": "new_ticket",
                "ticket_id": ticket_id,
                "ticket_no": ticket_no,
                "title": ticket.title,
                "message": f"{user['real_name']}提交了新工单"
            }, it_user['id'])
    
    conn.commit()
    
    # 发送webhook通知给相关人员（提交人和被分配的IT人员）
    webhook_title, webhook_content = get_webhook_message('ticket_created', {
        "ticket_no": ticket_no,
        "title": ticket.title,
        "submitter": user['real_name'],
        "priority": ticket.priority,
        "time": now_text(conn)
    }, conn)
    
    notify_user_ids = [payload['user_id']]  # 提交人
    if assigned_it:
        notify_user_ids.append(assigned_it['it_user_id'])  # 被分配的IT人员
    
    await send_webhook_notification('ticket_created', webhook_title, webhook_content, notify_user_ids, conn)
    log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '创建工单', f"创建工单 {ticket_no}：{ticket.title}", request)
    conn.commit()
    
    return {"message": "工单创建成功", "ticket_id": ticket_id, "ticket_no": ticket_no}

@app.get("/api/tickets/{ticket_id}")
async def get_ticket_detail(ticket_id: int, token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取工单详情
    cursor.execute("""
        SELECT t.*, c.category_name, 
               submitter.real_name as submitter_real_name,
               submitter.department as submitter_department,
               assignee.real_name as assignee_name
        FROM tickets t
        JOIN ticket_categories c ON t.category_id = c.id
        JOIN users submitter ON t.submitter_id = submitter.id
        LEFT JOIN users assignee ON t.assigned_to = assignee.id
        WHERE t.id = %s
    """, (ticket_id,))
    ticket = cursor.fetchone()
    
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    
    # 获取工单日志
    cursor.execute("""
        SELECT * FROM ticket_logs 
        WHERE ticket_id = %s 
        ORDER BY created_at DESC
    """, (ticket_id,))
    logs = cursor.fetchall()
    
    ticket['logs'] = logs
    workflow = parse_workflow_snapshot(ticket.get('workflow_snapshot'))
    ticket['workflow'] = workflow
    ticket['can_approve_workflow'] = False
    ticket['current_workflow_node'] = None
    if workflow and ticket.get('status') == 'pending_approval':
        current_index = workflow.get('current_index', 0)
        if 0 <= current_index < len(workflow.get('nodes', [])):
            current_node = workflow['nodes'][current_index]
            ticket['current_workflow_node'] = current_node
            ticket['can_approve_workflow'] = payload['user_id'] in normalize_id_list(current_node.get('approver_ids'))
    
    return {"ticket": ticket}

@app.delete("/api/tickets/{ticket_id}")
async def delete_ticket(ticket_id: int, token: str, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT t.*, c.category_name, u.username, u.real_name, r.role_type
        FROM tickets t
        JOIN ticket_categories c ON t.category_id = c.id
        JOIN users u ON u.id = %s
        JOIN roles r ON u.role_id = r.id
        WHERE t.id = %s
    """, (payload['user_id'], ticket_id))
    row = cursor.fetchone()

    if not row:
        raise HTTPException(status_code=404, detail="工单不存在")

    role_type = row['role_type']
    can_delete = role_type in ['super_admin', 'admin']
    if role_type == 'it':
        can_delete = can_user_see_ticket(row, payload['user_id'])

    if not can_delete:
        raise HTTPException(status_code=403, detail="您没有权限删除该工单")

    ticket_no = row['ticket_no']
    title = row['title']
    cursor.execute("DELETE FROM ticket_logs WHERE ticket_id = %s", (ticket_id,))
    cursor.execute("DELETE FROM notifications WHERE ticket_id = %s", (ticket_id,))
    cursor.execute("DELETE FROM tickets WHERE id = %s", (ticket_id,))
    log_operation(conn, payload['user_id'], row['username'], row['real_name'], 'ticket', '删除工单', f"删除工单 {ticket_no}：{title}", request)
    conn.commit()

    return {"message": "工单已删除"}

@app.put("/api/tickets/{ticket_id}/workflow")
async def approve_ticket_workflow(ticket_id: int, token: str, action: WorkflowAction, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")

    action_name = (action.action or "").strip().lower()
    if action_name not in ["approve", "reject"]:
        raise HTTPException(status_code=400, detail="审批操作只能是通过或驳回")

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT t.*, submitter.real_name as submitter_real_name
        FROM tickets t
        JOIN users submitter ON t.submitter_id = submitter.id
        WHERE t.id = %s
    """, (ticket_id,))
    ticket = cursor.fetchone()
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    if ticket['status'] != 'pending_approval':
        raise HTTPException(status_code=400, detail="该工单当前不在审批中")

    workflow = parse_workflow_snapshot(ticket.get('workflow_snapshot'))
    if not workflow:
        raise HTTPException(status_code=400, detail="该工单没有审批流程")

    current_index = workflow.get('current_index', 0)
    nodes = workflow.get('nodes', [])
    if current_index < 0 or current_index >= len(nodes):
        raise HTTPException(status_code=400, detail="审批流程状态异常")

    current_node = nodes[current_index]
    if payload['user_id'] not in normalize_id_list(current_node.get('approver_ids')):
        raise HTTPException(status_code=403, detail="您不是当前节点审批人")

    cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    comment = (action.comment or "").strip()
    action_text = "通过" if action_name == "approve" else "驳回"

    current_node["status"] = "approved" if action_name == "approve" else "rejected"
    current_node["approved_by"] = payload['user_id']
    current_node["approved_by_name"] = user['real_name']
    current_node["approved_at"] = now_text(conn)
    current_node["approval_comment"] = comment

    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], f"审批{action_text}", f"{user['real_name']}在“{current_node.get('name')}”节点{action_text}了工单" + (f"，意见：{comment}" if comment else "")))

    if action_name == "reject":
        workflow["status"] = "rejected"
        workflow["current_index"] = current_index
        cursor.execute("""
            UPDATE tickets SET status = 'rejected', workflow_snapshot = %s WHERE id = %s
        """, (json.dumps(workflow, ensure_ascii=False), ticket_id))
        cursor.execute("""
            INSERT INTO notifications (user_id, ticket_id, title, content, type)
            VALUES (%s, %s, %s, %s, %s)
        """, (ticket['submitter_id'], ticket_id, '工单审批被驳回', f"您的工单：{ticket['title']} 在“{current_node.get('name')}”节点被驳回", 'workflow_rejected'))
        await manager.send_personal_message({
            "type": "workflow_rejected",
            "ticket_id": ticket_id,
            "ticket_no": ticket['ticket_no'],
            "title": ticket['title'],
            "message": f"您的工单在“{current_node.get('name')}”节点被驳回"
        }, ticket['submitter_id'])
        log_operation(conn, user['id'], user['username'], user['real_name'], 'workflow', '驳回工单', f"驳回工单 {ticket['ticket_no']}：{ticket['title']}", request)
        conn.commit()
        return {"message": "已驳回"}

    next_index = current_index + 1
    if next_index < len(nodes):
        nodes[next_index]["status"] = "pending"
        workflow["current_index"] = next_index
        workflow["status"] = "pending"
        next_node = nodes[next_index]
        cursor.execute("UPDATE tickets SET workflow_snapshot = %s WHERE id = %s", (json.dumps(workflow, ensure_ascii=False), ticket_id))
        for approver_id in normalize_id_list(next_node.get("approver_ids")):
            cursor.execute("""
                INSERT INTO notifications (user_id, ticket_id, title, content, type)
                VALUES (%s, %s, %s, %s, %s)
            """, (approver_id, ticket_id, '待审批工单', f"工单：{ticket['title']} 已流转到“{next_node.get('name')}”节点，请您审批", 'workflow_notice'))
            await manager.send_personal_message({
                "type": "workflow_notice",
                "ticket_id": ticket_id,
                "ticket_no": ticket['ticket_no'],
                "title": ticket['title'],
                "message": f"工单已流转到“{next_node.get('name')}”节点，请您审批"
            }, approver_id)
        log_operation(conn, user['id'], user['username'], user['real_name'], 'workflow', '审批通过', f"审批通过工单 {ticket['ticket_no']}：{ticket['title']}", request)
        conn.commit()
        return {"message": "审批通过，已流转到下一节点"}

    workflow["status"] = "approved"
    workflow["current_index"] = current_index
    cursor.execute("UPDATE tickets SET workflow_snapshot = %s WHERE id = %s", (json.dumps(workflow, ensure_ascii=False), ticket_id))
    assigned_it = await assign_ticket_after_workflow(cursor, ticket_id, ticket, ticket['submitter_real_name'])
    cursor.execute("""
        INSERT INTO notifications (user_id, ticket_id, title, content, type)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket['submitter_id'], ticket_id, '工单审批通过', f"您的工单：{ticket['title']} 已全部审批通过", 'workflow_approved'))
    await manager.send_personal_message({
        "type": "workflow_approved",
        "ticket_id": ticket_id,
        "ticket_no": ticket['ticket_no'],
        "title": ticket['title'],
        "message": "您的工单已全部审批通过"
    }, ticket['submitter_id'])
    log_operation(conn, user['id'], user['username'], user['real_name'], 'workflow', '审批完成', f"审批完成工单 {ticket['ticket_no']}：{ticket['title']}", request)
    conn.commit()
    return {"message": "审批通过，流程已完成", "assigned_to": assigned_it['it_user_id'] if assigned_it else None}

@app.put("/api/tickets/{ticket_id}/claim")
async def claim_ticket(ticket_id: int, token: str, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 检查用户是否为IT人员
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    if user['role_type'] != 'it':
        raise HTTPException(status_code=403, detail="只有IT人员可以认领工单")
    
    # 获取工单信息
    cursor.execute("SELECT * FROM tickets WHERE id = %s", (ticket_id,))
    ticket = cursor.fetchone()
    
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    
    if ticket['status'] != 'pending':
        raise HTTPException(status_code=400, detail="该工单已被认领")
    
    # 更新工单状态
    cursor.execute("""
        UPDATE tickets 
        SET status = 'claimed', assigned_to = %s, claimed_at = NOW()
        WHERE id = %s
    """, (payload['user_id'], ticket_id))
    
    # 记录日志
    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], '认领工单', f"{user['real_name']}认领了工单"))
    
    # 通知提交人
    cursor.execute("""
        INSERT INTO notifications (user_id, ticket_id, title, content, type)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket['submitter_id'], ticket_id, '工单已被认领', f"您的工单已被{user['real_name']}认领", 'ticket_claimed'))
    
    # 发送WebSocket通知
    await manager.send_personal_message({
        "type": "ticket_claimed",
        "ticket_id": ticket_id,
        "message": f"您的工单已被{user['real_name']}认领"
    }, ticket['submitter_id'])
    
    conn.commit()

    webhook_title, webhook_content = get_webhook_message('ticket_claimed', {
        "ticket_no": ticket['ticket_no'],
        "title": ticket['title'],
        "operator": user['real_name'],
        "time": now_text(conn)
    }, conn)
    await send_webhook_notification('ticket_claimed', webhook_title, webhook_content, [ticket['submitter_id'], payload['user_id']], conn)
    log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '认领工单', f"认领工单 {ticket['ticket_no']}：{ticket['title']}", request)
    conn.commit()
    
    return {"message": "认领成功"}

@app.put("/api/tickets/{ticket_id}/process")
async def process_ticket(ticket_id: int, token: str, request: Request, message: Optional[str] = None, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取工单信息
    cursor.execute("SELECT * FROM tickets WHERE id = %s", (ticket_id,))
    ticket = cursor.fetchone()
    
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    
    if ticket['assigned_to'] != payload['user_id']:
        raise HTTPException(status_code=403, detail="只能处理自己认领的工单")
    
    # 获取用户信息
    cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    # 更新工单状态
    cursor.execute("""
        UPDATE tickets 
        SET status = 'processing'
        WHERE id = %s
    """, (ticket_id,))
    
    # 记录日志
    log_content = f"{user['real_name']}开始处理工单"
    if message:
        log_content += f"，备注：{message}"
    
    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], '开始处理', log_content))
    
    # 通知提交人
    notification_content = f"{user['real_name']}正在处理您的工单"
    if message:
        notification_content += f"，备注：{message}"
    
    cursor.execute("""
        INSERT INTO notifications (user_id, ticket_id, title, content, type)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket['submitter_id'], ticket_id, '工单处理中', notification_content, 'ticket_processing'))
    
    # 发送WebSocket实时通知给提交人
    await manager.send_personal_message({
        "type": "ticket_processing",
        "ticket_id": ticket_id,
        "ticket_no": ticket['ticket_no'],
        "message": notification_content
    }, ticket['submitter_id'])
    
    conn.commit()

    webhook_title, webhook_content = get_webhook_message('ticket_processing', {
        "ticket_no": ticket['ticket_no'],
        "title": ticket['title'],
        "operator": user['real_name'],
        "message": message or '无',
        "time": now_text(conn)
    }, conn)
    await send_webhook_notification('ticket_processing', webhook_title, webhook_content, [ticket['submitter_id'], payload['user_id']], conn)
    log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '开始处理', f"开始处理工单 {ticket['ticket_no']}：{ticket['title']}", request)
    conn.commit()
    
    return {"message": "已开始处理"}

@app.put("/api/tickets/{ticket_id}/complete")
async def complete_ticket(ticket_id: int, token: str, request: Request, solution: Optional[str] = "", conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取工单信息
    cursor.execute("SELECT * FROM tickets WHERE id = %s", (ticket_id,))
    ticket = cursor.fetchone()
    
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    
    if ticket['assigned_to'] != payload['user_id']:
        raise HTTPException(status_code=403, detail="只能完成自己认领的工单")
    
    # 获取用户信息
    cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    solution_text = (solution or "").strip()

    # 更新工单状态
    cursor.execute("""
        UPDATE tickets 
        SET status = 'completed', solution = %s, completed_at = NOW()
        WHERE id = %s
    """, (solution_text, ticket_id))
    
    # 记录日志
    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], '完成工单', f"{user['real_name']}完成了工单，解决方案：{solution_text or '未填写'}"))
    
    # 通知提交人
    cursor.execute("""
        INSERT INTO notifications (user_id, ticket_id, title, content, type)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket['submitter_id'], ticket_id, '工单已完成', f"您的工单已完成，请确认", 'ticket_completed'))
    
    # 发送WebSocket通知
    await manager.send_personal_message({
        "type": "ticket_completed",
        "ticket_id": ticket_id,
        "message": "您的工单已完成，请确认"
    }, ticket['submitter_id'])
    
    conn.commit()
    
    # 发送webhook通知给相关人员（提交人和处理人）
    webhook_title, webhook_content = get_webhook_message('ticket_completed', {
        "ticket_no": ticket['ticket_no'],
        "title": ticket['title'],
        "operator": user['real_name'],
        "solution": solution_text or '未填写',
        "time": now_text(conn)
    }, conn)
    
    notify_user_ids = [ticket['submitter_id'], payload['user_id']]  # 提交人和处理人
    await send_webhook_notification('ticket_completed', webhook_title, webhook_content, notify_user_ids, conn)
    log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '完成工单', f"完成工单 {ticket['ticket_no']}：{ticket['title']}，解决方案：{solution_text or '未填写'}", request)
    conn.commit()
    
    return {"message": "工单已完成"}

@app.put("/api/tickets/{ticket_id}/rate")
async def rate_ticket(ticket_id: int, token: str, rating: TicketRate, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取工单信息
    cursor.execute("SELECT * FROM tickets WHERE id = %s", (ticket_id,))
    ticket = cursor.fetchone()
    
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    
    if ticket['submitter_id'] != payload['user_id']:
        raise HTTPException(status_code=403, detail="只能评价自己提交的工单")
        
    if ticket['status'] != 'completed':
        raise HTTPException(status_code=400, detail="只能评价已完成的工单")
    
    # 更新工单满意度
    cursor.execute("""
        UPDATE tickets 
        SET satisfaction = %s, satisfaction_comment = %s
        WHERE id = %s
    """, (rating.satisfaction, rating.comment, ticket_id))
    
    # 获取用户信息
    cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    # 记录日志
    log_content = f"{user['real_name']}评价了工单：{rating.satisfaction}星"
    if rating.comment:
        log_content += f"，评论：{rating.comment}"
        
    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], '评价工单', log_content))
    log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '评价工单', f"评价工单 {ticket['ticket_no']}：{rating.satisfaction}星", request)
    
    # 通知IT人员
    if ticket['assigned_to']:
        cursor.execute("""
            INSERT INTO notifications (user_id, ticket_id, title, content, type)
            VALUES (%s, %s, %s, %s, %s)
        """, (ticket['assigned_to'], ticket_id, '工单收到评价', f"用户评价了您的服务：{rating.satisfaction}星", 'ticket_rated'))
        
        # 发送WebSocket通知
        await manager.send_personal_message({
            "type": "ticket_rated",
            "ticket_id": ticket_id,
            "message": f"用户评价了您的服务：{rating.satisfaction}星"
        }, ticket['assigned_to'])
    
    conn.commit()
    
    return {"message": "评价提交成功"}

@app.put("/api/tickets/{ticket_id}/close")
async def close_ticket(ticket_id: int, token: str, request: Request, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取工单信息
    cursor.execute("SELECT * FROM tickets WHERE id = %s", (ticket_id,))
    ticket = cursor.fetchone()
    
    if not ticket:
        raise HTTPException(status_code=404, detail="工单不存在")
    
    if ticket['submitter_id'] != payload['user_id']:
        # 检查是否为管理员
        cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
        user = cursor.fetchone()
        if user['role_type'] not in ['super_admin', 'admin']:
            raise HTTPException(status_code=403, detail="只能结单自己提交的工单")
    else:
        cursor.execute("SELECT * FROM users WHERE id = %s", (payload['user_id'],))
        user = cursor.fetchone()
    
    # 更新工单状态
    cursor.execute("""
        UPDATE tickets 
        SET status = 'closed', closed_at = NOW()
        WHERE id = %s
    """, (ticket_id,))
    
    # 记录日志
    cursor.execute("""
        INSERT INTO ticket_logs (ticket_id, user_id, user_name, action, content)
        VALUES (%s, %s, %s, %s, %s)
    """, (ticket_id, payload['user_id'], user['real_name'], '结单', f"{user['real_name']}结束了工单"))
    
    # 通知IT人员
    if ticket['assigned_to']:
        cursor.execute("""
            INSERT INTO notifications (user_id, ticket_id, title, content, type)
            VALUES (%s, %s, %s, %s, %s)
        """, (ticket['assigned_to'], ticket_id, '工单已结单', f"工单已被用户确认结单", 'ticket_closed'))
        
        # 发送WebSocket通知
        await manager.send_personal_message({
            "type": "ticket_closed",
            "ticket_id": ticket_id,
            "message": "工单已被用户确认结单"
        }, ticket['assigned_to'])
    
    conn.commit()

    webhook_title, webhook_content = get_webhook_message('ticket_closed', {
        "ticket_no": ticket['ticket_no'],
        "title": ticket['title'],
        "operator": user['real_name'],
        "time": now_text(conn)
    }, conn)
    notify_user_ids = [ticket['submitter_id']]
    if ticket['assigned_to']:
        notify_user_ids.append(ticket['assigned_to'])
    await send_webhook_notification('ticket_closed', webhook_title, webhook_content, notify_user_ids, conn)
    log_operation(conn, user['id'], user['username'], user['real_name'], 'ticket', '结单', f"结单工单 {ticket['ticket_no']}：{ticket['title']}", request)
    conn.commit()
    
    return {"message": "工单已结单"}

# 通知接口
@app.get("/api/audit/login-logs")
async def get_login_logs(token: str, limit: int = 100, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    limit = max(1, min(300, limit))
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT *
        FROM login_logs
        ORDER BY created_at DESC
        LIMIT %s
    """, (limit,))
    return {"logs": cursor.fetchall()}

@app.get("/api/audit/operation-logs")
async def get_operation_logs(token: str, limit: int = 100, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    limit = max(1, min(300, limit))
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT *
        FROM operation_logs
        ORDER BY created_at DESC
        LIMIT %s
    """, (limit,))
    return {"logs": cursor.fetchall()}

@app.get("/api/notifications")
async def get_notifications(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT * FROM notifications 
        WHERE user_id = %s 
        ORDER BY created_at DESC 
        LIMIT 50
    """, (payload['user_id'],))
    notifications = cursor.fetchall()
    
    return {"notifications": notifications}

@app.put("/api/notifications/{notification_id}/read")
async def mark_notification_read(notification_id: int, token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        UPDATE notifications 
        SET is_read = 1 
        WHERE id = %s AND user_id = %s
    """, (notification_id, payload['user_id']))
    
    conn.commit()
    
    return {"message": "已标记为已读"}

@app.put("/api/notifications/read-all")
async def mark_all_notifications_read(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        UPDATE notifications 
        SET is_read = 1 
        WHERE user_id = %s AND is_read = 0
    """, (payload['user_id'],))
    
    conn.commit()
    
    return {"message": "所有通知已标记为已读"}

# 企业消息推送功能
async def send_dingtalk_message(webhook_url: str, secret: Optional[str], title: str, content: str):
    """发送钉钉消息"""
    try:
        timestamp = str(round(time.time() * 1000))
        sign = None
        
        if secret:
            string_to_sign = f'{timestamp}\n{secret}'
            hmac_code = hmac.new(secret.encode('utf-8'), string_to_sign.encode('utf-8'), digestmod=hashlib.sha256).digest()
            sign = base64.b64encode(hmac_code).decode('utf-8')
        
        url = webhook_url
        if sign:
            url = f"{webhook_url}&timestamp={timestamp}&sign={sign}"
        
        message = {
            "msgtype": "markdown",
            "markdown": {
                "title": title,
                "text": f"### {title}\n\n{content}"
            }
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(url, json=message, timeout=10.0)
            return response.status_code == 200
    except Exception as e:
        print(f"钉钉消息发送失败: {e}")
        return False

async def send_feishu_message(webhook_url: str, secret: Optional[str], title: str, content: str):
    """发送飞书消息"""
    try:
        timestamp = str(int(time.time()))
        sign = None
        
        if secret:
            string_to_sign = f'{timestamp}\n{secret}'
            hmac_code = hmac.new(string_to_sign.encode('utf-8'), digestmod=hashlib.sha256).digest()
            sign = base64.b64encode(hmac_code).decode('utf-8')
        
        message = {
            "msg_type": "interactive",
            "card": {
                "header": {
                    "title": {
                        "tag": "plain_text",
                        "content": title
                    },
                    "template": "blue"
                },
                "elements": [
                    {
                        "tag": "div",
                        "text": {
                            "tag": "lark_md",
                            "content": content
                        }
                    }
                ]
            }
        }
        
        if sign:
            message["timestamp"] = timestamp
            message["sign"] = sign
        
        async with httpx.AsyncClient() as client:
            response = await client.post(webhook_url, json=message, timeout=10.0)
            return response.status_code == 200
    except Exception as e:
        print(f"飞书消息发送失败: {e}")
        return False

async def send_wechat_message(webhook_url: str, title: str, content: str):
    """发送企业微信消息"""
    try:
        message = {
            "msgtype": "markdown",
            "markdown": {
                "content": f"**{title}**\n\n{content}"
            }
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(webhook_url, json=message, timeout=10.0)
            return response.status_code == 200
    except Exception as e:
        print(f"企业微信消息发送失败: {e}")
        return False

def get_webhook_message(event_type: str, variables: dict, conn):
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("""
        SELECT title_template, content_template
        FROM webhook_templates
        WHERE event_type = %s AND is_active = 1
        ORDER BY updated_at DESC, id DESC
        LIMIT 1
    """, (event_type,))
    template = cursor.fetchone()

    if not template:
        template = DEFAULT_WEBHOOK_TEMPLATES.get(event_type, {
            "title": "IT运维工单系统通知",
            "content": "{title}"
        })
        title_template = template['title']
        content_template = template['content']
    else:
        title_template = template['title_template']
        content_template = template['content_template']

    return (
        render_text_template(title_template, variables),
        render_text_template(content_template, variables)
    )

async def send_webhook_notification(event_type: str, title: str, content: str, user_ids: List[int], conn):
    """发送webhook通知到指定用户的配置"""
    try:
        cursor = conn.cursor(pymysql.cursors.DictCursor)

        # 查询相关用户的个人配置，以及管理员设置的全局配置。
        if user_ids:
            unique_user_ids = list(dict.fromkeys(user_ids))
            placeholders = ','.join(['%s'] * len(unique_user_ids))
            cursor.execute(f"""
                SELECT * FROM webhook_configs
                WHERE enabled = 1
                  AND (
                    scope = 'global'
                    OR ((scope IS NULL OR scope = 'personal') AND user_id IN ({placeholders}))
                  )
            """, unique_user_ids)
        else:
            cursor.execute("""
                SELECT * FROM webhook_configs
                WHERE enabled = 1 AND scope = 'global'
            """)
        configs = cursor.fetchall()
        seen_config_ids = set()
        
        for config in configs:
            if config['id'] in seen_config_ids:
                continue
            seen_config_ids.add(config['id'])

            # 检查是否配置了事件过滤
            if config['notify_events']:
                try:
                    events = json.loads(config['notify_events'])
                    if event_type not in events:
                        continue
                except:
                    pass
            
            # 根据平台类型发送消息
            if config['platform'] == 'dingtalk':
                await send_dingtalk_message(config['webhook_url'], config['secret'], title, content)
            elif config['platform'] == 'feishu':
                await send_feishu_message(config['webhook_url'], config['secret'], title, content)
            elif config['platform'] == 'wechat':
                await send_wechat_message(config['webhook_url'], title, content)
    except Exception as e:
        print(f"Webhook通知发送失败: {e}")

# Webhook配置管理接口
@app.get("/api/webhook-configs")
async def get_webhook_configs(token: str, conn=Depends(get_db)):
    """获取webhook配置列表"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取当前用户信息
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    # 管理员可以看到所有配置，普通用户只能看到自己的配置
    if user['role_type'] in ['super_admin', 'admin']:
        cursor.execute("""
            SELECT w.*, u.real_name as user_name, u.username
            FROM webhook_configs w
            JOIN users u ON w.user_id = u.id
            ORDER BY w.created_at DESC
        """)
    else:
        cursor.execute("""
            SELECT w.*, u.real_name as user_name, u.username
            FROM webhook_configs w
            JOIN users u ON w.user_id = u.id
            WHERE w.user_id = %s
            ORDER BY w.created_at DESC
        """, (payload['user_id'],))
    
    configs = cursor.fetchall()
    
    return {"configs": configs}

@app.post("/api/webhook-configs")
async def create_webhook_config(token: str, config: WebhookConfigCreate, conn=Depends(get_db)):
    """创建webhook配置"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 验证平台类型
    if config.platform not in ['dingtalk', 'feishu', 'wechat']:
        raise HTTPException(status_code=400, detail="不支持的平台类型")
    
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()

    scope = config.scope or "personal"
    if scope not in ["personal", "global"]:
        raise HTTPException(status_code=400, detail="不支持的配置范围")
    if scope == "global" and user['role_type'] not in ['super_admin', 'admin']:
        raise HTTPException(status_code=403, detail="只有管理员可以创建全局推送配置")

    # 转换事件列表为JSON
    notify_events_json = json.dumps(config.notify_events) if config.notify_events else None
    
    # 创建配置，关联到当前用户
    cursor.execute("""
        INSERT INTO webhook_configs (user_id, name, platform, webhook_url, secret, scope, notify_events)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
    """, (payload['user_id'], config.name, config.platform, config.webhook_url, config.secret, scope, notify_events_json))
    
    conn.commit()
    
    return {"message": "配置创建成功", "config_id": cursor.lastrowid}

@app.put("/api/webhook-configs/{config_id}")
async def update_webhook_config(config_id: int, token: str, config: WebhookConfigUpdate, conn=Depends(get_db)):
    """更新webhook配置"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取当前用户信息
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    # 检查配置是否存在
    cursor.execute("SELECT * FROM webhook_configs WHERE id = %s", (config_id,))
    existing_config = cursor.fetchone()
    
    if not existing_config:
        raise HTTPException(status_code=404, detail="配置不存在")
    
    # 检查权限：只能修改自己的配置，管理员可以修改所有配置
    if user['role_type'] not in ['super_admin', 'admin'] and existing_config['user_id'] != payload['user_id']:
        raise HTTPException(status_code=403, detail="没有权限修改此配置")
    
    # 构建更新语句
    updates = []
    params = []
    
    if config.name is not None:
        updates.append("name = %s")
        params.append(config.name)
    
    if config.webhook_url is not None:
        updates.append("webhook_url = %s")
        params.append(config.webhook_url)
    
    if config.secret is not None:
        updates.append("secret = %s")
        params.append(config.secret)
    
    if config.enabled is not None:
        updates.append("enabled = %s")
        params.append(config.enabled)

    if config.scope is not None:
        if config.scope not in ["personal", "global"]:
            raise HTTPException(status_code=400, detail="不支持的配置范围")
        if config.scope == "global" and user['role_type'] not in ['super_admin', 'admin']:
            raise HTTPException(status_code=403, detail="只有管理员可以设置全局推送配置")
        updates.append("scope = %s")
        params.append(config.scope)
    
    if config.notify_events is not None:
        updates.append("notify_events = %s")
        params.append(json.dumps(config.notify_events))
    
    if not updates:
        raise HTTPException(status_code=400, detail="没有要更新的字段")
    
    params.append(config_id)
    sql = f"UPDATE webhook_configs SET {', '.join(updates)} WHERE id = %s"
    
    cursor.execute(sql, params)
    conn.commit()
    
    return {"message": "配置更新成功"}

@app.delete("/api/webhook-configs/{config_id}")
async def delete_webhook_config(config_id: int, token: str, conn=Depends(get_db)):
    """删除webhook配置"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取当前用户信息
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    # 检查配置是否存在
    cursor.execute("SELECT * FROM webhook_configs WHERE id = %s", (config_id,))
    existing_config = cursor.fetchone()
    
    if not existing_config:
        raise HTTPException(status_code=404, detail="配置不存在")
    
    # 检查权限：只能删除自己的配置，管理员可以删除所有配置
    if user['role_type'] not in ['super_admin', 'admin'] and existing_config['user_id'] != payload['user_id']:
        raise HTTPException(status_code=403, detail="没有权限删除此配置")
    
    cursor.execute("DELETE FROM webhook_configs WHERE id = %s", (config_id,))
    conn.commit()
    
    return {"message": "配置删除成功"}

@app.post("/api/webhook-configs/{config_id}/test")
async def test_webhook_config(config_id: int, token: str, conn=Depends(get_db)):
    """测试webhook配置"""
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取当前用户信息
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    user = cursor.fetchone()
    
    # 获取配置
    cursor.execute("SELECT * FROM webhook_configs WHERE id = %s", (config_id,))
    config = cursor.fetchone()
    
    if not config:
        raise HTTPException(status_code=404, detail="配置不存在")
    
    # 检查权限：只能测试自己的配置，管理员可以测试所有配置
    if user['role_type'] not in ['super_admin', 'admin'] and config['user_id'] != payload['user_id']:
        raise HTTPException(status_code=403, detail="没有权限测试此配置")
    
    # 发送测试消息
    title = "IT运维工单系统 - 测试消息"
    content = f"这是一条测试消息\n\n发送时间：{now_text(conn)}\n\n如果您收到此消息，说明webhook配置正确！"
    
    success = False
    if config['platform'] == 'dingtalk':
        success = await send_dingtalk_message(config['webhook_url'], config['secret'], title, content)
    elif config['platform'] == 'feishu':
        success = await send_feishu_message(config['webhook_url'], config['secret'], title, content)
    elif config['platform'] == 'wechat':
        success = await send_wechat_message(config['webhook_url'], title, content)
    
    if success:
        return {"message": "测试消息发送成功"}
    else:
        raise HTTPException(status_code=500, detail="测试消息发送失败")

@app.get("/api/settings/ticket-number")
async def get_ticket_number_settings(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    return {
        "prefix": get_setting(cursor, 'ticket_no_prefix', 'TK'),
        "date_format": get_setting(cursor, 'ticket_no_date_format', '%Y%m%d%H%M%S'),
        "random_digits": int(get_setting(cursor, 'ticket_no_random_digits', '3'))
    }

@app.put("/api/settings/ticket-number")
async def update_ticket_number_settings(token: str, settings: TicketNumberSettings, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    prefix = settings.prefix.strip()
    if not prefix or len(prefix) > 20:
        raise HTTPException(status_code=400, detail="编号前缀不能为空且不能超过20个字符")
    if settings.random_digits < 1 or settings.random_digits > 12:
        raise HTTPException(status_code=400, detail="随机码位数需在1到12之间")

    try:
        datetime.now().strftime(settings.date_format)
    except Exception:
        raise HTTPException(status_code=400, detail="日期格式不正确")

    cursor = conn.cursor()
    rows = [
        ('ticket_no_prefix', prefix),
        ('ticket_no_date_format', settings.date_format),
        ('ticket_no_random_digits', str(settings.random_digits))
    ]
    for key, value in rows:
        cursor.execute("""
            INSERT INTO system_settings (setting_key, setting_value)
            VALUES (%s, %s)
            ON DUPLICATE KEY UPDATE setting_value = VALUES(setting_value)
        """, (key, value))
    conn.commit()

    return {"message": "工单编号规则已保存"}

@app.get("/api/webhook-templates")
async def get_webhook_templates(token: str, event_type: Optional[str] = None, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    if event_type:
        cursor.execute("""
            SELECT t.*, u.real_name AS creator_name
            FROM webhook_templates t
            LEFT JOIN users u ON t.created_by = u.id
            WHERE t.event_type = %s
            ORDER BY t.event_type, t.is_active DESC, t.updated_at DESC
        """, (event_type,))
    else:
        cursor.execute("""
            SELECT t.*, u.real_name AS creator_name
            FROM webhook_templates t
            LEFT JOIN users u ON t.created_by = u.id
            ORDER BY t.event_type, t.is_active DESC, t.updated_at DESC
        """)

    return {
        "templates": cursor.fetchall(),
        "events": [
            {"value": "ticket_created", "label": "工单创建"},
            {"value": "ticket_claimed", "label": "工单认领"},
            {"value": "ticket_processing", "label": "工单处理中"},
            {"value": "ticket_completed", "label": "工单完成"},
            {"value": "ticket_closed", "label": "工单结单"}
        ],
        "variables": ["ticket_no", "title", "submitter", "operator", "priority", "message", "solution", "time"]
    }

@app.post("/api/webhook-templates")
async def create_webhook_template(token: str, template: WebhookTemplateCreate, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    if template.event_type not in DEFAULT_WEBHOOK_TEMPLATES:
        raise HTTPException(status_code=400, detail="不支持的事件类型")

    cursor = conn.cursor()
    if template.is_active:
        cursor.execute("UPDATE webhook_templates SET is_active = 0 WHERE event_type = %s", (template.event_type,))
    cursor.execute("""
        INSERT INTO webhook_templates
            (name, event_type, title_template, content_template, is_active, created_by)
        VALUES (%s, %s, %s, %s, %s, %s)
    """, (
        template.name,
        template.event_type,
        template.title_template,
        template.content_template,
        1 if template.is_active else 0,
        payload['user_id']
    ))
    conn.commit()

    return {"message": "模板创建成功", "template_id": cursor.lastrowid}

@app.put("/api/webhook-templates/{template_id}")
async def update_webhook_template(template_id: int, token: str, template: WebhookTemplateUpdate, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM webhook_templates WHERE id = %s", (template_id,))
    existing = cursor.fetchone()
    if not existing:
        raise HTTPException(status_code=404, detail="模板不存在")

    event_type = template.event_type or existing['event_type']
    if event_type not in DEFAULT_WEBHOOK_TEMPLATES:
        raise HTTPException(status_code=400, detail="不支持的事件类型")

    if template.is_active:
        cursor.execute("UPDATE webhook_templates SET is_active = 0 WHERE event_type = %s AND id <> %s", (event_type, template_id))

    updates = []
    params = []
    for field, value in [
        ("name", template.name),
        ("event_type", template.event_type),
        ("title_template", template.title_template),
        ("content_template", template.content_template),
        ("is_active", template.is_active)
    ]:
        if value is not None:
            updates.append(f"{field} = %s")
            params.append(value)

    if not updates:
        raise HTTPException(status_code=400, detail="没有要更新的字段")

    params.append(template_id)
    cursor.execute(f"UPDATE webhook_templates SET {', '.join(updates)} WHERE id = %s", params)
    conn.commit()

    return {"message": "模板更新成功"}

@app.delete("/api/webhook-templates/{template_id}")
async def delete_webhook_template(template_id: int, token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    require_admin(payload, conn)

    cursor = conn.cursor(pymysql.cursors.DictCursor)
    cursor.execute("SELECT * FROM webhook_templates WHERE id = %s", (template_id,))
    template = cursor.fetchone()
    if not template:
        raise HTTPException(status_code=404, detail="模板不存在")

    if template['is_active']:
        raise HTTPException(status_code=400, detail="启用中的模板不能删除，请先启用同事件下其他模板")

    cursor.execute("DELETE FROM webhook_templates WHERE id = %s", (template_id,))
    conn.commit()
    return {"message": "模板删除成功"}

# WebSocket连接
@app.websocket("/ws/{token}")
async def websocket_endpoint(websocket: WebSocket, token: str):
    payload = verify_token(token)
    if not payload:
        await websocket.close(code=1008)
        return
    
    user_id = payload['user_id']
    await manager.connect(websocket, user_id)
    
    try:
        while True:
            data = await websocket.receive_text()
            # 处理客户端消息（如果需要）
    except WebSocketDisconnect:
        manager.disconnect(user_id)

# 统计接口
@app.get("/api/statistics")
async def get_statistics(token: str, conn=Depends(get_db)):
    payload = verify_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="无效的token")
    
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # 获取当前用户信息
    cursor.execute("SELECT u.*, r.role_type FROM users u JOIN roles r ON u.role_id = r.id WHERE u.id = %s", (payload['user_id'],))
    current_user = cursor.fetchone()
    
    stats = {}
    
    if current_user['role_type'] in ['super_admin', 'admin', 'hr', 'it']:
        # 管理员、HR、IT运维可以看到所有工单统计
        cursor.execute("SELECT COUNT(*) as total FROM tickets")
        stats['total_tickets'] = cursor.fetchone()['total']
        
        # 待处理工单数
        cursor.execute("SELECT COUNT(*) as pending FROM tickets WHERE status IN ('pending', 'pending_approval')")
        stats['pending_tickets'] = cursor.fetchone()['pending']
        
        # 处理中工单数
        cursor.execute("SELECT COUNT(*) as processing FROM tickets WHERE status IN ('claimed', 'processing')")
        stats['processing_tickets'] = cursor.fetchone()['processing']
        
        # 已完成工单数（包括completed和closed状态）
        cursor.execute("SELECT COUNT(*) as completed FROM tickets WHERE status IN ('completed', 'closed')")
        stats['completed_tickets'] = cursor.fetchone()['completed']
        
        # 已结单工单数
        cursor.execute("SELECT COUNT(*) as closed FROM tickets WHERE status = 'closed'")
        stats['closed_tickets'] = cursor.fetchone()['closed']
        
        # 平均满意度
        cursor.execute("SELECT AVG(satisfaction) as avg_satisfaction FROM tickets WHERE satisfaction IS NOT NULL")
        result = cursor.fetchone()
        stats['avg_satisfaction'] = round(result['avg_satisfaction'], 1) if result and result['avg_satisfaction'] else 0
        
        # 如果是IT人员，额外显示个人的平均满意度
        if current_user['role_type'] == 'it':
            cursor.execute("SELECT AVG(satisfaction) as my_avg FROM tickets WHERE assigned_to = %s AND satisfaction IS NOT NULL", (payload['user_id'],))
            result = cursor.fetchone()
            stats['my_avg_satisfaction'] = round(result['my_avg'], 1) if result and result['my_avg'] else 0
            
        # 如果是 super_admin, admin, hr，显示每个工程师的平均满意度
        if current_user['role_type'] in ['super_admin', 'admin', 'hr']:
            cursor.execute("""
                SELECT u.id, u.real_name, u.username, AVG(t.satisfaction) as avg_satisfaction, COUNT(t.satisfaction) as rated_count
                FROM users u
                JOIN roles r ON u.role_id = r.id
                LEFT JOIN tickets t ON u.id = t.assigned_to AND t.satisfaction IS NOT NULL
                WHERE r.role_type = 'it'
                GROUP BY u.id
                ORDER BY avg_satisfaction DESC
            """)
            engineer_stats = cursor.fetchall()
            stats['engineer_stats'] = [
                {
                    'id': row['id'],
                    'name': row['real_name'],
                    'username': row['username'],
                    'avg_satisfaction': round(row['avg_satisfaction'], 1) if row['avg_satisfaction'] else 0,
                    'rated_count': row['rated_count']
                }
                for row in engineer_stats
            ]
            
    else:
        # 普通用户只能看到自己的工单统计
        cursor.execute("SELECT COUNT(*) as total FROM tickets WHERE submitter_id = %s", (payload['user_id'],))
        stats['my_tickets'] = cursor.fetchone()['total']
        
        # 待处理
        cursor.execute("SELECT COUNT(*) as pending FROM tickets WHERE submitter_id = %s AND status IN ('pending', 'pending_approval')", (payload['user_id'],))
        stats['pending_tickets'] = cursor.fetchone()['pending']
        
        # 处理中
        cursor.execute("SELECT COUNT(*) as processing FROM tickets WHERE submitter_id = %s AND status IN ('claimed', 'processing')", (payload['user_id'],))
        stats['processing_tickets'] = cursor.fetchone()['processing']
        
        # 已完成（包括completed和closed状态）
        cursor.execute("SELECT COUNT(*) as completed FROM tickets WHERE submitter_id = %s AND status IN ('completed', 'closed')", (payload['user_id'],))
        stats['completed_tickets'] = cursor.fetchone()['completed']
        
        # 个人平均满意度（自己给出的评价）
        cursor.execute("SELECT AVG(satisfaction) as avg_satisfaction FROM tickets WHERE submitter_id = %s AND satisfaction IS NOT NULL", (payload['user_id'],))
        result = cursor.fetchone()
        stats['avg_satisfaction'] = round(result['avg_satisfaction'], 1) if result and result['avg_satisfaction'] else 0
    
    return stats

# 挂载静态文件（必须在最后）
app.mount("/static", StaticFiles(directory="static", html=True), name="static")
